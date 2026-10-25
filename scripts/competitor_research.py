#!/usr/bin/env python3
"""Free competitor research for every published page (no paid keys).
For each page: search (ddgs, Bing backend, in-en) -> read top competitor pages -> rule-based
extraction of headings, FAQ questions, rule facts (₹, %, dates, sections) and gaps vs our page.
Writes research_out/competitor_research_<date>.xlsx (uploaded as a workflow artifact, never committed)."""
import os, re, json, time, datetime, pathlib, html
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = pathlib.Path(os.environ.get('SITE_DIR', pathlib.Path(__file__).resolve().parent.parent))
OUT = pathlib.Path(os.environ.get('OUT_DIR', ROOT / 'research_out')); OUT.mkdir(parents=True, exist_ok=True)
SEARXNG = os.environ.get('SEARXNG_URL', '').rstrip('/')
BUDGET_MIN = int(os.environ.get('BUDGET_MIN', '200'))
MAX_COMP = int(os.environ.get('MAX_COMP', '5'))
MY = 'udyoggrowth.com'
UA = {'User-Agent': 'Mozilla/5.0 (compatible; UdyogResearch/1.0; +https://udyoggrowth.com)'}
SKIP = ('assets', 'sitemap', 'docs', 'scripts', 'seo', 'reports', 'data', 'images', 'dist', 'private', 'pages', 'tools', 'research_out', 'lead-backend')
FACT_RE = re.compile(r'(₹|Rs\.?\s?\d|\d+\s?%|\b(19|20)\d{2}\b|\bsection\s+\d|\bform\s+[A-Z0-9-]+|\bby\s+\d{1,2}(st|nd|rd|th)?\b|\b\d+\s+(days|months|years)\b)', re.I)
DEADLINE = time.time() + BUDGET_MIN * 60

def our_pages():
    out = []
    for f in sorted(ROOT.glob('**/index.html')):
        rel = f.parent.relative_to(ROOT).as_posix()
        if rel == '.' or rel.split('/')[0] in SKIP: continue
        h = f.read_text(encoding='utf-8', errors='ignore')
        cl = lambda x: html.unescape(re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', x))).strip() if x else ''
        t = re.search(r'<title>(.*?)</title>', h, re.S); h1 = re.search(r'<h1[^>]*>(.*?)</h1>', h, re.S)
        body = re.sub(r'<script.*?</script>|<style.*?</style>', '', h, flags=re.S)
        out.append({'slug': rel, 'url': f'https://{MY}/{rel}/', 'title': cl(t.group(1)) if t else '',
                    'h1': cl(h1.group(1)) if h1 else '', 'text': cl(body).lower()})
    return out

def search(q):
    if SEARXNG:
        for attempt in range(3):
            try:
                r = requests.get(SEARXNG + '/search', params={'q': q, 'format': 'json', 'language': 'en-IN'}, timeout=40)
                if r.status_code == 200:
                    return [x.get('url') for x in r.json().get('results', []) if x.get('url')][:8]
            except (requests.RequestException, ValueError):
                pass
            time.sleep(5 * (attempt + 1))
        return []
    try:
        from ddgs import DDGS
        res = DDGS(timeout=15).text(q, region='in-en', backend='bing', max_results=8) or []
        return [r.get('href') or r.get('url') for r in res if (r.get('href') or r.get('url'))]
    except Exception:
        return []

def read(url):
    try:
        r = requests.get(url, headers=UA, timeout=25)
        if r.status_code != 200 or 'html' not in r.headers.get('content-type', ''): return None
        s = BeautifulSoup(r.text, 'html.parser')
        for x in s(['script', 'style', 'nav', 'footer', 'header', 'noscript', 'form']): x.decompose()
        heads = [h.get_text(' ', strip=True) for h in s.find_all(['h1', 'h2', 'h3'])][:80]
        faqs = [h for h in heads if h.endswith('?')]
        faqs += [d.get_text(' ', strip=True) for d in s.find_all('summary')][:20]
        sents = re.split(r'(?<=[.!?])\s+', re.sub(r'\s+', ' ', s.get_text(' ', strip=True)))
        facts = [x for x in sents if FACT_RE.search(x) and 40 < len(x) < 300][:25]
        return {'heads': heads, 'faqs': list(dict.fromkeys(faqs))[:15], 'facts': facts}
    except requests.RequestException:
        return None

def keywords(s):
    return {w for w in re.findall(r'[a-z]{4,}', s.lower()) if w not in {'with', 'that', 'this', 'from', 'your', 'india', 'delhi', 'consultant', 'service', 'services', 'their', 'what', 'when', 'will', 'have'}}

def gaps(comp_heads, our_text):
    ours = keywords(our_text)
    out = []
    for h in comp_heads:
        k = keywords(h)
        if k and len(k & ours) / len(k) < 0.5:
            out.append(h)
    return out[:12]

def main():
    rows = []; progress = []
    pages = our_pages()
    print('pages', len(pages), flush=True)
    for p in pages:
        if time.time() > DEADLINE: print('budget reached', flush=True); break
        q = (p['h1'] or p['title']).replace('—', ' ')[:120] + ' India'
        links = [u for u in search(q) if u and MY not in u][:MAX_COMP]
        n = 0
        for u in links:
            if time.time() > DEADLINE: break
            c = read(u)
            time.sleep(2)
            if not c: continue
            n += 1
            rows.append([p['slug'], p['url'], u, ' | '.join(c['heads'][:25]), ' | '.join(c['faqs']),
                         '\n'.join(c['facts'][:12]), ' | '.join(gaps(c['heads'], p['text']))])
        progress.append([p['slug'], q, n, datetime.datetime.now().isoformat(timespec='seconds')])
        print('done', p['slug'], 'competitors', n, flush=True)
    write(rows, progress)

def write(rows, progress):
    wb = Workbook(); ws = wb.active; ws.title = 'Competitor research'
    hdr = ['Our page', 'Our URL', 'Competitor URL', 'Competitor headings', 'FAQ questions covered', 'Rule facts (with figures/dates)', 'Gaps for our page (not covered by our text)']
    ws.append(hdr)
    for c in ws[1]: c.fill = PatternFill('solid', fgColor='1F3A5F'); c.font = Font(bold=True, color='FFFFFF')
    for r in rows: ws.append(r)
    for i, w in enumerate([34, 44, 50, 70, 60, 80, 60], 1): ws.column_dimensions[chr(64 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions
    s = wb.create_sheet('Progress'); s.append(['Our page', 'Search query', 'Competitor pages read', 'Finished at'])
    for r in progress: s.append(r)
    name = OUT / f'competitor_research_{datetime.date.today().isoformat()}.xlsx'
    wb.save(name); print('saved', name, 'rows', len(rows), flush=True)

if __name__ == '__main__':
    main()
